from datetime import datetime, timezone
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from sqlalchemy import select, func, delete
from sqlalchemy.orm import Session
import os
import csv
import subprocess
import tempfile
import mimetypes
from pathlib import Path
from jose import JWTError, jwt
from google.oauth2 import id_token
from google.auth.transport import requests as google_requests
import bcrypt

from app.config import settings
from app.db.database import init_db, get_session, SessionLocal, Faq, PendingQuestion, Intent, Document, Feedback, User, Tenant, Invite
from app.schemas import (
    IngestRequest,
    IngestResponse,
    QueryRequest,
    QueryResponse,
    QueryHit,
    StatsResponse,
    FaqCreate,
    FaqUpdate,
    FaqStatusUpdate,
    FaqOut,
    PendingOut,
    PendingLog,
    IntentCreate,
    IntentOut,
    IntentUtterance,
    MergePreview,
    MergeMatch,
    MergeRequest,
    MergeResult,
    FaqSuggestion,
    PendingSuggestion,
    IntentSuggestions,
    FaqMergePreview,
    FaqMergeRequest,
    FaqMergeResult,
    DocumentInfo,
    DocumentDeleteResponse,
    FeedbackCreate,
    FeedbackOut,
    GoogleAuthRequest,
    TokenResponse,
    UserOut,
    TenantOut,
    InviteCreate,
    InviteOut,
    SystemAdminLogin,
    AcceptInviteRequest,
)
from app.services.ingestion import ingest_text
from app.services import vector_store, nlp, apps
from app.services.nlp import chunk_text


app = FastAPI(title="Assistance RAG API", version="0.2.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=settings.cors_list,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def on_startup():
    init_db()
    _ensure_systemadmin()
    _resync_approved_faqs_to_chroma()
    os.makedirs(settings.upload_dir, exist_ok=True)


pwd_context = None


def _hash_pin(pin: str) -> str:
    salt = bcrypt.gensalt()
    hashed = bcrypt.hashpw(pin.encode("utf-8"), salt)
    return hashed.decode("utf-8")


def _ensure_systemadmin() -> None:
    with SessionLocal() as db:
        user = db.execute(select(User).where(User.is_systemadmin == True)).scalar_one_or_none()
        if user:
            if settings.systemadmin_pin:
                user.pin_hash = _hash_pin(settings.systemadmin_pin)
                db.commit()
            return

        pin = settings.systemadmin_pin if settings.systemadmin_pin else "123456"
        hashed_pin = _hash_pin(pin)
        
        user = User(
            email=settings.systemadmin_email,
            name="System Admin",
            is_systemadmin=True,
            pin_hash=hashed_pin,
        )
        db.add(user)
        db.commit()


def _create_access_token(data: dict) -> str:
    to_encode = data.copy()
    expire = datetime.now(timezone.utc).timestamp() + settings.jwt_expire_minutes * 60
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, settings.jwt_secret_key, algorithm=settings.jwt_algorithm)


def _get_bearer_user(authorization: str = Header(""), db: Session = Depends(get_session)) -> User:
    if not authorization.startswith("Bearer "):
        raise HTTPException(401, "Missing bearer token")
    token = authorization.split(" ", 1)[1]
    try:
        payload = jwt.decode(token, settings.jwt_secret_key, algorithms=[settings.jwt_algorithm])
        user_id = int(payload.get("sub"))
    except (JWTError, ValueError, TypeError):
        raise HTTPException(401, "Invalid token")

    user = db.get(User, user_id)
    if not user:
        raise HTTPException(401, "User not found")
    return user


def _extract_text_from_file(path: str, content_type: str | None = None) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".txt":
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    if ext == ".csv":
        with open(path, "r", encoding="utf-8", errors="ignore", newline="") as f:
            reader = csv.reader(f)
            return "\n".join(" | ".join(row) for row in reader)
    if ext in (".xlsx", ".xls"):
        import pandas as pd
        df = pd.read_excel(path)
        return "\n".join(" | ".join(map(str, row)) for row in df.values.tolist())
    if ext == ".docx":
        from docx import Document as DocxDocument
        doc = DocxDocument(path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if ext == ".doc":
        with tempfile.TemporaryDirectory() as tmpdir:
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "docx", "--outdir", tmpdir, path],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )
            converted = Path(tmpdir) / (Path(path).stem + ".docx")
            if not converted.exists():
                raise HTTPException(500, "Failed to convert DOC file")
            from docx import Document as DocxDocument
            doc = DocxDocument(str(converted))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    raise HTTPException(415, f"Unsupported file type: {ext}")


def _convert_for_preview(path: Path) -> Path | None:
    ext = path.suffix.lower()
    if ext in ('.docx', '.doc', '.xlsx', '.xls', '.csv', '.txt'):
        with tempfile.TemporaryDirectory() as tmpdir:
            subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", tmpdir, str(path)],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
            )
            converted = Path(tmpdir) / (path.stem + ".pdf")
            if converted.exists():
                return converted
    return None


def _resync_approved_faqs_to_chroma() -> None:
    """On startup, make sure every 'approved' FAQ is in the Chroma faq collection.

    Postgres is the source of truth; Chroma is rebuilt from it so that we recover
    after wiping .chroma or moving to a new machine.
    """
    with SessionLocal() as db:
        approved = db.execute(
            select(Faq).where(Faq.status == "approved")
        ).scalars().all()
        for faq in approved:
            vector_store.upsert_faq(faq.id, faq.question, faq.answer, faq.status)

        pending = db.execute(
            select(PendingQuestion).where(PendingQuestion.resolved == False)  # noqa: E712
        ).scalars().all()
        for p in pending:
            vector_store.index_pending(p.id, p.question)

        intents = db.execute(select(Intent)).scalars().all()
        for intent in intents:
            vector_store.upsert_intent(
                intent.id, intent.name, intent.response_template, intent.faq_id
            )


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/ingest", response_model=IngestResponse)
def ingest(req: IngestRequest, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot ingest documents")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    return ingest_text(db, current_user.tenant_id, req.source, req.text)


@app.post("/upload")
def upload_file(file: UploadFile = File(...), authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot upload documents")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")

    ext = Path(file.filename).suffix.lower()
    allowed = {".txt", ".csv", ".xlsx", ".xls", ".docx", ".doc", ".pdf"}
    if ext not in allowed:
        raise HTTPException(415, f"Unsupported file type: {ext}")

    upload_path = Path(settings.upload_dir)
    upload_path.mkdir(parents=True, exist_ok=True)
    dest = upload_path / file.filename
    with dest.open("wb") as out:
        out.write(file.file.read())

    text = _extract_text_from_file(str(dest), file.content_type)
    chunks = chunk_text(text)
    if not chunks:
        return {"source": file.filename, "chunks": 0, "status": "uploaded"}

    start_id = vector_store.collection_size()
    ids = [f"{dest.name}-{start_id + i}" for i in range(len(chunks))]
    vector_store.add_chunks(dest.name, chunks, ids, current_user.tenant_id)

    for i, chunk in enumerate(chunks):
        db.add(Document(source=dest.name, chunk_index=start_id + i, text=chunk, tenant_id=current_user.tenant_id))
    db.commit()

    return {"source": dest.name, "chunks": len(chunks), "status": "ingested"}


@app.get("/documents", response_model=list[DocumentInfo])
def list_documents(authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        rows = db.execute(
            select(Document.source, func.count(Document.id), func.min(Document.created_at))
            .group_by(Document.source)
            .order_by(Document.source)
        ).all()
    else:
        rows = db.execute(
            select(Document.source, func.count(Document.id), func.min(Document.created_at))
            .where(Document.tenant_id == current_user.tenant_id)
            .group_by(Document.source)
            .order_by(Document.source)
        ).all()
    return [
        DocumentInfo(source=source, chunks=count, created_at=created_at)
        for source, count, created_at in rows
    ]


@app.delete("/documents/{filename}", response_model=DocumentDeleteResponse)
def delete_document(filename: str, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    source = filename
    if current_user.is_systemadmin:
        rows = db.execute(select(func.count(Document.id)).where(Document.source == source)).scalar_one()
        db.execute(delete(Document).where(Document.source == source))
    else:
        rows = db.execute(
            select(func.count(Document.id)).where(Document.source == source, Document.tenant_id == current_user.tenant_id)
        ).scalar_one()
        db.execute(delete(Document).where(Document.source == source, Document.tenant_id == current_user.tenant_id))
    db.commit()
    vector_store.delete_chunks_by_source(source)
    upload_path = Path(settings.upload_dir) / filename
    if upload_path.exists():
        upload_path.unlink()
    return DocumentDeleteResponse(source=source, deleted_chunks=rows, status="deleted")


@app.get("/documents/{filename}/view")
def view_document(filename: str):
    upload_path = Path(settings.upload_dir) / filename
    if not upload_path.exists():
        raise HTTPException(404, "File not found")

    ext = upload_path.suffix.lower()
    if ext == '.pdf':
        return FileResponse(
            str(upload_path),
            filename=filename,
            media_type='application/pdf',
            headers={'Content-Disposition': 'inline'},
        )

    preview_path = _convert_for_preview(upload_path)
    if preview_path:
        return FileResponse(
            str(preview_path),
            filename=filename,
            media_type='application/pdf',
            headers={'Content-Disposition': 'inline'},
        )

    return FileResponse(
        str(upload_path),
        filename=filename,
        media_type='text/plain',
        headers={'Content-Disposition': 'inline'},
    )


@app.get("/documents/{filename}/download")
def download_document(filename: str):
    upload_path = Path(settings.upload_dir) / filename
    if not upload_path.exists():
        raise HTTPException(404, "File not found")
    return FileResponse(
        str(upload_path),
        filename=filename,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f"attachment; filename=\"{filename}\""},
    )


@app.post("/feedback", response_model=FeedbackOut)
def create_feedback(req: FeedbackCreate, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot create feedback")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    feedback = Feedback(
        message_id=req.message_id,
        question=req.question,
        answer=req.answer,
        source=req.source,
        rating=req.rating,
        tenant_id=current_user.tenant_id,
    )
    db.add(feedback)
    db.commit()
    db.refresh(feedback)
    return feedback


@app.post("/auth/google", response_model=TokenResponse)
def auth_google(req: GoogleAuthRequest, db: Session = Depends(get_session)):
    try:
        payload = id_token.verify_token(
            req.id_token,
            google_requests.Request(),
            audience=None,
        )
    except ValueError:
        raise HTTPException(401, "Invalid Google ID token")

    google_id = payload.get("sub")
    email = payload.get("email")
    name = payload.get("name") or email
    picture = payload.get("picture")

    if not google_id or not email:
        raise HTTPException(401, "Invalid Google ID token payload")

    user = db.execute(select(User).where(User.google_id == google_id)).scalar_one_or_none()
    if not user:
        user = db.execute(select(User).where(User.email == email)).scalar_one_or_none()

    invite = None
    if req.invite_token:
        invite = db.execute(select(Invite).where(Invite.token == req.invite_token)).scalar_one_or_none()
        if not invite:
            raise HTTPException(404, "Invite not found")
        if invite.used_at is not None:
            raise HTTPException(400, "Invite already used")
        if datetime.utcnow() > invite.expires_at:
            raise HTTPException(400, "Invite expired")

    if user:
        user.name = name
        user.picture = picture
        if invite and not user.tenant_id:
            user.tenant_id = invite.tenant_id
            user.role = invite.role
            invite.used_at = datetime.utcnow()
            db.commit()
    else:
        if invite:
            user = User(
                google_id=google_id,
                email=email,
                name=name,
                picture=picture,
                tenant_id=invite.tenant_id,
                role=invite.role,
            )
            invite.used_at = datetime.utcnow()
        else:
            tenant = Tenant(name=name or email, slug=email.split("@")[0])
            db.add(tenant)
            db.flush()
            user = User(
                google_id=google_id,
                email=email,
                name=name,
                picture=picture,
                tenant_id=tenant.id,
                role="superadmin",
            )
        db.add(user)

    db.commit()
    db.refresh(user)
    token = _create_access_token({
        "sub": str(user.id),
        "email": user.email,
        "tenant_id": user.tenant_id,
        "role": user.role,
        "is_systemadmin": user.is_systemadmin,
    })
    return TokenResponse(access_token=token)


@app.get("/auth/systemadmin/status")
def systemadmin_status(db: Session = Depends(get_session)):
    user = db.execute(select(User).where(User.is_systemadmin == True)).scalar_one_or_none()
    if not user:
        return {"configured": False}
    return {
        "configured": True,
        "email": user.email,
    }


@app.post("/auth/systemadmin", response_model=TokenResponse)
def auth_systemadmin(req: SystemAdminLogin, db: Session = Depends(get_session)):
    if not settings.systemadmin_pin:
        raise HTTPException(401, "System admin PIN not configured")

    user = db.execute(
        select(User).where(User.is_systemadmin == True)
    ).scalar_one_or_none()
    if not user:
        raise HTTPException(401, "System admin not configured")

    if not req.pin or not bcrypt.checkpw(req.pin.encode("utf-8"), user.pin_hash.encode("utf-8")):
        raise HTTPException(401, "Invalid PIN")

    token = _create_access_token({
        "sub": str(user.id),
        "email": user.email,
        "tenant_id": user.tenant_id,
        "role": user.role,
        "is_systemadmin": True,
    })
    return TokenResponse(access_token=token)


@app.post("/auth/accept-invite", response_model=TokenResponse)
def auth_accept_invite(req: AcceptInviteRequest, invite_token: str, db: Session = Depends(get_session)):
    invite = db.execute(select(Invite).where(Invite.token == invite_token)).scalar_one_or_none()
    if not invite:
        raise HTTPException(404, "Invite not found")
    if invite.used_at is not None:
        raise HTTPException(400, "Invite already used")
    if datetime.utcnow() > invite.expires_at:
        raise HTTPException(400, "Invite expired")

    try:
        payload = id_token.verify_token(
            req.id_token,
            google_requests.Request(),
            audience=None,
        )
    except ValueError:
        raise HTTPException(401, "Invalid Google ID token")

    google_id = payload.get("sub")
    email = payload.get("email")
    name = payload.get("name") or email
    picture = payload.get("picture")

    if not google_id or not email:
        raise HTTPException(401, "Invalid Google ID token payload")

    user = db.execute(select(User).where(User.google_id == google_id)).scalar_one_or_none()
    if not user:
        user = db.execute(select(User).where(User.email == email)).scalar_one_or_none()

    if user:
        user.name = name
        user.picture = picture
        if not user.tenant_id:
            user.tenant_id = invite.tenant_id
            user.role = invite.role
    else:
        user = User(
            google_id=google_id,
            email=email,
            name=name,
            picture=picture,
            tenant_id=invite.tenant_id,
            role=invite.role,
        )
        db.add(user)

    invite.used_at = datetime.utcnow()
    db.commit()
    db.refresh(user)
    token = _create_access_token({
        "sub": str(user.id),
        "email": user.email,
        "tenant_id": user.tenant_id,
        "role": user.role,
        "is_systemadmin": user.is_systemadmin,
    })
    return TokenResponse(access_token=token)


@app.post("/auth/invites", response_model=InviteOut)
def create_invite(req: InviteCreate, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot create invites")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    if current_user.role not in ("superadmin", "admin"):
        raise HTTPException(403, "Only admins can create invites")

    token_val = secrets.token_urlsafe(32)
    expires_at = datetime.utcnow() + __import__("datetime").timedelta(hours=settings.invite_expire_hours)
    invite = Invite(
        tenant_id=current_user.tenant_id,
        email=req.email,
        token=token_val,
        role=req.role,
        created_by=current_user.id,
        expires_at=expires_at,
    )
    db.add(invite)
    db.commit()
    db.refresh(invite)
    return invite


@app.get("/auth/invites", response_model=list[InviteOut])
def list_invites(authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot list tenant invites")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    if current_user.role not in ("superadmin", "admin"):
        raise HTTPException(403, "Only admins can view invites")

    invites = db.execute(
        select(Invite).where(Invite.tenant_id == current_user.tenant_id).order_by(Invite.created_at.desc())
    ).scalars().all()
    return invites


@app.get("/auth/tenants", response_model=list[TenantOut])
def list_tenants(authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if not current_user.is_systemadmin:
        raise HTTPException(403, "Only system admin can list tenants")

    tenants = db.execute(select(Tenant).order_by(Tenant.created_at.desc())).scalars().all()
    return tenants


@app.get("/auth/me", response_model=UserOut)
def auth_me(authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    return current_user


def _default_response(db: Session, req: QueryRequest, ents: list, disabled_external_sources: list[str] | None, tenant_id: int | None = None) -> QueryResponse:
    pending = _record_pending(db, req.question, tenant_id)
    return QueryResponse(
        question=req.question,
        source="pending",
        answer="I don't have an answer for that yet. A trainer has been notified.",
        hits=[],
        entities=ents,
        pending_id=pending.id,
        disabled_external_sources=disabled_external_sources,
    )


@app.post("/query", response_model=QueryResponse)
def query(req: QueryRequest, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot query")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")

    tenant_id = current_user.tenant_id
    ents = nlp.extract_entities(req.question)
    apps_cfg = req.apps or {}
    location = req.location or {}
    detected = apps.detect_intents(req.question)
    disabled_external_sources: list[str] = []
    if "weather" in detected and not apps_cfg.get("weather"):
        disabled_external_sources.append("weather")
    if "wiki" in detected and not apps_cfg.get("wiki"):
        disabled_external_sources.append("wiki")

    faq_hits = vector_store.search_faq(req.question, n_results=1, tenant_id=tenant_id)
    if faq_hits:
        best = faq_hits[0]
        if best["distance"] is not None and best["distance"] <= settings.faq_match_threshold:
            return QueryResponse(
                question=req.question,
                source="faq",
                answer=best["answer"],
                hits=[QueryHit(text=best["question"], metadata={"source": "faq", "faq_id": best["faq_id"]}, distance=best["distance"])],
                entities=ents,
                disabled_external_sources=disabled_external_sources or None,
            )
        if best["distance"] is not None and best["distance"] > 1:
            return _default_response(db, req, ents, disabled_external_sources or None, tenant_id)

    loose = vector_store.search_faq_loose(req.question, n_results=3, tenant_id=tenant_id)
    intent_hit: dict | None = None
    if loose:
        candidate = loose[0]
        d = candidate.get("distance")
        if (
            d is not None
            and settings.faq_match_threshold < d <= settings.intent_match_threshold
        ):
            intent_hit = {
                "matched_question": candidate.get("question"),
                "matched_faq_id": candidate.get("faq_id"),
                "distance": d,
                "answer": candidate.get("answer"),
            }

    if intent_hit is not None:
        return QueryResponse(
                question=req.question,
                source="intent",
                answer=intent_hit["answer"],
                hits=[
                    QueryHit(
                        text=intent_hit["matched_question"],
                        metadata={
                            "source": "faq",
                            "faq_id": intent_hit["matched_faq_id"],
                            "via": "intent",
                        },
                        distance=intent_hit["distance"],
                    )
                ],
                entities=ents,
                intent=intent_hit,
                disabled_external_sources=disabled_external_sources or None,
            )

    doc_hits = vector_store.query(req.question, n_results=req.n_results, tenant_id=tenant_id)
    if doc_hits:
        if doc_hits[0].get("distance") is not None and doc_hits[0]["distance"] > 1:
            return _default_response(db, req, ents, disabled_external_sources or None, tenant_id)
        return QueryResponse(
                question=req.question,
                source="docs",
                answer=doc_hits[0]["text"],
                hits=[QueryHit(**h) for h in doc_hits],
                entities=ents,
                disabled_external_sources=disabled_external_sources or None,
            )

    if apps_cfg.get("weather") and "weather" in detected:
        lat = location.get("lat")
        lon = location.get("lon")
        result = apps.weather_for(req.question, lat=lat, lon=lon)
        if result:
            return QueryResponse(
                question=req.question,
                source="weather",
                answer=result["answer"],
                hits=[],
                entities=ents,
                app=result["payload"],
                disabled_external_sources=disabled_external_sources or None,
            )
    if apps_cfg.get("wiki") and "wiki" in detected:
        result = apps.wiki_summary(req.question)
        if result:
            return QueryResponse(
                question=req.question,
                source="wiki",
                answer=result["answer"],
                hits=[
                    QueryHit(
                        text=result["payload"].get("title", req.question),
                        metadata={
                            "source": "wiki",
                            "url": result["payload"].get("url"),
                        },
                        distance=None,
                    )
                ],
                entities=ents,
                app=result["payload"],
                disabled_external_sources=disabled_external_sources or None,
            )

    if disabled_external_sources:
        parts = []
        if "weather" in disabled_external_sources:
            parts.append("Weather")
        if "wiki" in disabled_external_sources:
            parts.append("Wikipedia")
        joined = " and ".join(parts)
        return QueryResponse(
            question=req.question,
            source="pending",
            answer=f"I don't have a trained answer for that yet. Enable {joined} in Settings to get live answers from your location or the web.",
            hits=[],
            entities=ents,
            pending_id=None,
            disabled_external_sources=disabled_external_sources or None,
        )

    pending = _record_pending(db, req.question, tenant_id)
    return QueryResponse(
        question=req.question,
        source="pending",
        answer="I don't have an answer for that yet. A trainer has been notified.",
        hits=[],
        entities=ents,
        pending_id=pending.id,
        disabled_external_sources=disabled_external_sources or None,
    )


@app.post("/faq", response_model=FaqOut)
def create_faq(req: FaqCreate, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot create FAQs")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    faq = Faq(
        question=req.question.strip(),
        answer=req.answer.strip(),
        status="pending",
        tenant_id=current_user.tenant_id,
    )
    db.add(faq)
    db.commit()
    db.refresh(faq)
    vector_store.upsert_faq(faq.id, faq.question, faq.answer, faq.status, current_user.tenant_id)
    return faq


@app.get("/faq", response_model=list[FaqOut])
def list_faqs(status: str | None = None, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    stmt = select(Faq).order_by(Faq.updated_at.desc())
    if not current_user.is_systemadmin:
        stmt = stmt.where(Faq.tenant_id == current_user.tenant_id)
    if status:
        stmt = stmt.where(Faq.status == status)
    return db.execute(stmt).scalars().all()


@app.patch("/faq/{faq_id}", response_model=FaqOut)
def update_faq(faq_id: int, req: FaqUpdate, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot update FAQs")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    faq = db.get(Faq, faq_id)
    if not faq or faq.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "FAQ not found")
    faq.answer = req.answer.strip()
    db.commit()
    db.refresh(faq)
    vector_store.upsert_faq(faq.id, faq.question, faq.answer, faq.status, faq.tenant_id)
    return faq


@app.patch("/faq/{faq_id}/status", response_model=FaqOut)
def update_faq_status(faq_id: int, req: FaqStatusUpdate, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot update FAQ status")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    allowed = {"pending", "approved", "declined", "duplicate"}
    if req.status not in allowed:
        raise HTTPException(400, f"status must be one of {sorted(allowed)}")
    faq = db.get(Faq, faq_id)
    if not faq or faq.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "FAQ not found")
    faq.status = req.status
    db.commit()
    db.refresh(faq)
    if faq.status == "approved":
        vector_store.upsert_faq(faq.id, faq.question, faq.answer, faq.status, faq.tenant_id)
    else:
        vector_store.delete_faq(faq.id)
    return faq


@app.delete("/faq/{faq_id}")
def remove_faq(faq_id: int, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot delete FAQs")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    faq = db.get(Faq, faq_id)
    if not faq or faq.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "FAQ not found")
    db.delete(faq)
    db.commit()
    vector_store.delete_faq(faq_id)
    return {"deleted": faq_id}


@app.get("/faq/inbox", response_model=list[PendingOut])
def list_inbox(resolved: bool = False, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot view inbox")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    return db.execute(
        select(PendingQuestion)
        .where(PendingQuestion.resolved == resolved, PendingQuestion.tenant_id == current_user.tenant_id)
        .order_by(PendingQuestion.occurrences.desc(), PendingQuestion.last_seen_at.desc())
    ).scalars().all()


@app.post("/faq/inbox/{pending_id}/answer", response_model=FaqOut)
def answer_pending(pending_id: int, req: FaqUpdate, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot answer pending questions")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    pending = db.get(PendingQuestion, pending_id)
    if not pending or pending.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "Pending question not found")
    faq = Faq(question=pending.question, answer=req.answer.strip(), status="pending", tenant_id=current_user.tenant_id)
    db.add(faq)
    pending.resolved = True
    db.commit()
    db.refresh(faq)
    vector_store.upsert_faq(faq.id, faq.question, faq.answer, faq.status, current_user.tenant_id)
    vector_store.delete_pending_index(pending_id)
    return faq


@app.post("/faq/log", response_model=PendingOut)
def log_question(req: PendingLog, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot log questions")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    return _record_pending(db, req.question, current_user.tenant_id)


@app.delete("/faq/inbox/{pending_id}")
def delete_pending(pending_id: int, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot delete pending questions")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    pending = db.get(PendingQuestion, pending_id)
    if not pending or pending.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "Pending question not found")
    db.delete(pending)
    db.commit()
    vector_store.delete_pending_index(pending_id)
    return {"deleted": pending_id}


@app.get("/faq/pending/{pending_id}/merge-preview", response_model=MergePreview)
def merge_preview(pending_id: int, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot view merge preview")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    canonical = db.get(PendingQuestion, pending_id)
    if not canonical or canonical.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "Pending question not found")
    candidates = db.execute(
        select(PendingQuestion).where(
            PendingQuestion.resolved == False,
            PendingQuestion.id != canonical.id,
            PendingQuestion.tenant_id == current_user.tenant_id,
        )
    ).scalars().all()
    candidate_texts = [c.question for c in candidates]
    if not candidate_texts:
        return MergePreview(
            canonical_id=canonical.id,
            canonical_question=canonical.question,
            answer="",
            matches=[],
        )
    hits = vector_store.search_pending(
        canonical.question, candidate_texts, n_results=max(50, len(candidate_texts)), tenant_id=current_user.tenant_id
    )
    by_question = {h["question"]: h for h in hits}
    matches: list[MergeMatch] = []
    for c in candidates:
        h = by_question.get(c.question)
        if not h or h["distance"] is None:
            continue
        if h["distance"] <= settings.intent_merge_threshold:
            matches.append(
                MergeMatch(
                    pending_id=c.id,
                    question=c.question,
                    occurrences=c.occurrences,
                    distance=h["distance"],
                )
            )
    matches.sort(key=lambda m: m.distance)
    return MergePreview(
        canonical_id=canonical.id,
        canonical_question=canonical.question,
        answer="",
        matches=matches,
    )


@app.post("/faq/pending/{pending_id}/merge", response_model=MergeResult)
def merge_pending(
    pending_id: int,
    req: MergeRequest,
    authorization: str = Header(""),
    db: Session = Depends(get_session),
):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot merge pending questions")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    canonical = db.get(PendingQuestion, pending_id)
    if not canonical or canonical.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "Pending question not found")
    answer = req.answer.strip()
    if not answer:
        raise HTTPException(400, "answer is required")

    targets = []
    if req.pending_ids:
        for tid in req.pending_ids:
            t = db.get(PendingQuestion, tid)
            if t and not t.resolved and t.id != canonical.id and t.tenant_id == current_user.tenant_id:
                targets.append(t)

    faq = Faq(
        question=canonical.question.strip(),
        answer=answer,
        status="pending",
        tenant_id=current_user.tenant_id,
    )
    db.add(faq)
    db.commit()
    db.refresh(faq)
    vector_store.upsert_faq(faq.id, faq.question, faq.answer, faq.status, current_user.tenant_id)

    created = 1
    for t in targets:
        alias = Faq(
            question=t.question.strip(),
            answer=answer,
            status="pending",
            tenant_id=current_user.tenant_id,
        )
        db.add(alias)
        db.flush()
        vector_store.upsert_faq(alias.id, alias.question, alias.answer, faq.status, current_user.tenant_id)
        t.resolved = True
        vector_store.delete_pending_index(t.id)
        created += 1

    canonical.resolved = True
    vector_store.delete_pending_index(canonical.id)
    db.commit()
    return MergeResult(
        canonical_faq_id=faq.id,
        created=created,
        resolved=len(targets) + 1,
    )


@app.get("/faq/intent-suggestions", response_model=IntentSuggestions)
def intent_suggestions(question: str, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot view intent suggestions")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    q = question.strip()
    faq_hits = vector_store.search_faq_loose(q, n_results=10, tenant_id=current_user.tenant_id)
    pending_hits = vector_store.search_pending_all(q, n_results=10, tenant_id=current_user.tenant_id)
    return IntentSuggestions(
        query=q,
        faqs=[
            FaqSuggestion(
                faq_id=h["faq_id"],
                question=h["question"],
                answer=h.get("answer", ""),
                status=h.get("status") or "unknown",
                distance=h.get("distance") if h.get("distance") is not None else 1.0,
            )
            for h in faq_hits
        ],
        pending=[
            PendingSuggestion(
                pending_id=h["pending_id"],
                question=h["question"],
                occurrences=next(
                    (p.occurrences for p in db.execute(
                        select(PendingQuestion).where(PendingQuestion.id == h["pending_id"])
                    ).scalars() if p),
                    1,
                ),
                distance=h.get("distance") if h.get("distance") is not None else 1.0,
            )
            for h in pending_hits
        ],
    )


@app.get("/faq/{faq_id}/merge-preview", response_model=FaqMergePreview)
def faq_merge_preview(faq_id: int, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot view FAQ merge preview")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    canonical = db.get(Faq, faq_id)
    if not canonical or canonical.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "FAQ not found")
    others = db.execute(
        select(Faq).where(Faq.id != canonical.id, Faq.tenant_id == current_user.tenant_id)
    ).scalars().all()
    other_ids = [o.id for o in others]
    if not other_ids:
        return FaqMergePreview(
            canonical_id=canonical.id,
            canonical_question=canonical.question,
            canonical_answer=canonical.answer,
            matches=[],
        )
    hits = vector_store.search_faq_in(canonical.question, other_ids, n_results=max(50, len(other_ids)))
    matches: list[FaqSuggestion] = []
    for h in hits:
        if h.get("distance") is None or h["distance"] > settings.intent_merge_threshold:
            continue
        matches.append(
            FaqSuggestion(
                faq_id=h["faq_id"],
                question=h["question"],
                answer=h.get("answer", ""),
                status=h.get("status") or "unknown",
                distance=h["distance"],
            )
        )
    matches.sort(key=lambda m: m.distance)
    return FaqMergePreview(
        canonical_id=canonical.id,
        canonical_question=canonical.question,
        canonical_answer=canonical.answer,
        matches=matches,
    )


@app.post("/faq/{faq_id}/merge", response_model=FaqMergeResult)
def faq_merge(
    faq_id: int,
    req: FaqMergeRequest,
    authorization: str = Header(""),
    db: Session = Depends(get_session),
):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot merge FAQs")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    canonical = db.get(Faq, faq_id)
    if not canonical or canonical.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "FAQ not found")
    absorbed = 0
    recreated = 0

    target_id = req.target_id if req.target_id else canonical.id
    target = db.get(Faq, target_id)
    if not target or target.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "Target FAQ not found")

    if req.answer is not None:
        new_answer = req.answer.strip()
        if not new_answer:
            raise HTTPException(400, "answer cannot be empty")
        if target.id == canonical.id:
            target.answer = new_answer
        else:
            canonical.answer = new_answer
            db.flush()

    db.flush()
    target_id_final = target.id
    target_obj = db.get(Faq, target_id_final)
    final_answer = target_obj.answer

    absorbed_questions: list[str] = []
    for fid in req.faq_ids:
        other = db.get(Faq, fid)
        if not other or other.id == target_id_final or other.tenant_id != current_user.tenant_id:
            continue
        absorbed_questions.append(other.question)
        db.delete(other)
        vector_store.delete_faq(other.id)
        absorbed += 1

    db.commit()
    db.refresh(target_obj)
    vector_store.upsert_faq(target_obj.id, target_obj.question, target_obj.answer, target_obj.status)

    for q in absorbed_questions:
        alias = Faq(
            question=q.strip(),
            answer=final_answer,
            status="pending",
            tenant_id=current_user.tenant_id,
        )
        db.add(alias)
        db.flush()
        vector_store.upsert_faq(alias.id, alias.question, alias.answer, alias.status)
        recreated += 1
    db.commit()

    return FaqMergeResult(
        canonical_faq_id=target_id_final,
        absorbed=absorbed,
        recreated=recreated,
    )


@app.post("/intents", response_model=IntentOut)
def create_intent(req: IntentCreate, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot create intents")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    intent = Intent(
        name=req.name.strip(),
        response_template=req.response_template.strip(),
        faq_id=req.faq_id,
        tenant_id=current_user.tenant_id,
    )
    db.add(intent)
    db.commit()
    db.refresh(intent)
    vector_store.upsert_intent(intent.id, intent.name, intent.response_template, intent.faq_id, current_user.tenant_id)
    return intent


@app.get("/intents", response_model=list[IntentOut])
def list_intents(authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    stmt = select(Intent).order_by(Intent.updated_at.desc())
    if not current_user.is_systemadmin:
        stmt = stmt.where(Intent.tenant_id == current_user.tenant_id)
    return db.execute(stmt).scalars().all()


@app.patch("/intents/{intent_id}", response_model=IntentOut)
def update_intent(intent_id: int, req: IntentCreate, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot update intents")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    intent = db.get(Intent, intent_id)
    if not intent or intent.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "Intent not found")
    intent.name = req.name.strip()
    intent.response_template = req.response_template.strip()
    intent.faq_id = req.faq_id
    db.commit()
    db.refresh(intent)
    vector_store.upsert_intent(intent.id, intent.name, intent.response_template, intent.faq_id, current_user.tenant_id)
    return intent


@app.delete("/intents/{intent_id}")
def remove_intent(intent_id: int, authorization: str = Header(""), db: Session = Depends(get_session)):
    current_user = _get_bearer_user(authorization, db)
    if current_user.is_systemadmin:
        raise HTTPException(403, "System admin cannot delete intents")
    if not current_user.tenant_id:
        raise HTTPException(403, "User is not assigned to a tenant")
    intent = db.get(Intent, intent_id)
    if not intent or intent.tenant_id != current_user.tenant_id:
        raise HTTPException(404, "Intent not found")
    db.delete(intent)
    db.commit()
    vector_store.delete_intent(intent_id)
    return {"deleted": intent_id}


@app.get("/stats", response_model=StatsResponse)
def stats(db: Session = Depends(get_session)):
    faq_count = vector_store.faq_size()
    pending_count = db.execute(
        select(func.count(PendingQuestion.id)).where(PendingQuestion.resolved == False)  # noqa: E712
    ).scalar_one()
    return StatsResponse(
        chunks=vector_store.collection_size(),
        faqs=faq_count,
        pending=pending_count,
    )


def _record_pending(db: Session, question: str, tenant_id: int | None = None) -> PendingQuestion:
    q = question.strip()
    existing = db.execute(
        select(PendingQuestion)
        .where(PendingQuestion.question == q, PendingQuestion.resolved == False)  # noqa: E712
    ).scalar_one_or_none()
    if existing:
        existing.occurrences += 1
        existing.last_seen_at = datetime.utcnow()
        db.commit()
        db.refresh(existing)
        vector_store.index_pending(existing.id, existing.question, tenant_id)
        return existing
    pending = PendingQuestion(question=q, tenant_id=tenant_id)
    db.add(pending)
    db.commit()
    db.refresh(pending)
    vector_store.index_pending(pending.id, pending.question, tenant_id)
    return pending